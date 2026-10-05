"""Conversion rules independent of the desktop interface. No network operations."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .parsing import flexible_date, flexible_amount, amount_parts

MAX_ROWS = 100_000
MAX_COLUMNS = 256
DEFAULT_PROFILE = Path(__file__).parent / "profiles" / "demo.json"


@dataclass(frozen=True)
class UnsafeCell:
    reason: str


@dataclass
class Source:
    path: Path
    sha256: str
    sheets: dict[str, list[tuple]]


@dataclass
class Issue:
    severity: str
    row: int | None
    field: str
    message: str
    field_id: str | None = None


@dataclass
class Result:
    source: Source
    sheet: str
    header_row: int
    profile: dict
    mapping: dict[str, int | None]
    records: list[dict] = field(default_factory=list)
    source_rows: list[int] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    read_rows: int = 0
    blank_rows: list[int] = field(default_factory=list)
    invalid_rows: list[int] = field(default_factory=list)
    corrections: dict = field(default_factory=dict)
    changes: list[dict] = field(default_factory=list)

    @property
    def errors(self):
        return [issue for issue in self.issues if issue.severity == "error"]

def validate_profile(profile: dict) -> dict:
    if not isinstance(profile, dict) or profile.get("version") != 1:
        raise ValueError("Нужен JSON-профиль версии 1.")
    if not isinstance(profile.get("name"), str) or not profile["name"].strip():
        raise ValueError("Укажите название профиля.")
    sheet = profile.get("sheet_name", "")
    if not isinstance(sheet, str) or not sheet.strip() or len(sheet) > 31 or re.search(r"[\\/*?:\[\]]", sheet):
        raise ValueError("Некорректное название выходного листа.")
    fields = profile.get("fields")
    if not isinstance(fields, list) or not fields or len(fields) > MAX_COLUMNS:
        raise ValueError("Профиль должен содержать от 1 до 256 полей.")
    ids, titles = set(), set()
    for f in fields:
        if not isinstance(f, dict):
            raise ValueError("Каждое поле профиля должно быть объектом.")
        key, title = f.get("id"), f.get("title")
        if not isinstance(key, str) or not re.fullmatch(r"[a-z][a-z0-9_]*", key) or key in ids:
            raise ValueError("Идентификаторы полей должны быть уникальными: латиница, цифры, подчёркивание.")
        if not isinstance(title, str) or not title.strip() or title in titles or len(title) > 200:
            raise ValueError("Названия полей должны быть непустыми и уникальными (до 200 символов).")
        if f.get("kind") not in {"text", "date", "identifier", "amount", "currency"}:
            raise ValueError(f"Неизвестный тип поля: {key}.")
        if f.get('date_order', 'auto') not in {'auto', 'dmy', 'mdy'} or f.get('decimal_separator', 'auto') not in {'auto', '.', ','}:
            raise ValueError(f"Некорректные настройки распознавания: {key}.")
        if not isinstance(f.get('drop_time', False), bool):
            raise ValueError('drop_time должен быть true или false.')
        if not isinstance(f.get("required"), bool):
            raise ValueError(f"Укажите required: true или false для {key}.")
        aliases = f.get("aliases", [])
        if not isinstance(aliases, list) or any(not isinstance(a, str) for a in aliases):
            raise ValueError(f"aliases должен быть списком строк: {key}.")
        ids.add(key)
        titles.add(title)
    mapping = profile.get("mapping", {})
    if not isinstance(mapping, dict) or any(k not in ids or (v is not None and not isinstance(v, str)) for k, v in mapping.items()):
        raise ValueError("Сохранённое сопоставление должно содержать названия исходных столбцов или null.")
    return profile


def load_profile(path: str | Path = DEFAULT_PROFILE) -> dict:
    return validate_profile(json.loads(Path(path).read_text(encoding="utf-8-sig")))


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def normalize(value):
    return re.sub(r"[^\w]", "", str(value or "").casefold().replace("ё", "е"))


def load_source(path: str | Path) -> Source:
    path = Path(path).resolve()
    if path.suffix.lower() not in {".xlsx", ".csv"}:
        raise ValueError("Поддерживаются .xlsx и .csv. Старый .xls сохраните как .xlsx через Excel.")
    raw = path.read_bytes()
    sheets = {}
    if path.suffix.lower() == ".csv":
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV должен быть в UTF-8. Пересохраните файл с этой кодировкой.") from None
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(io.StringIO(content, newline=""), dialect):
            rows.append(tuple(row))
            if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS:
                raise ValueError("Лимит прототипа: 100 000 строк и 256 столбцов на лист.")
        sheets["CSV"] = rows
    else:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError("Лимит прототипа: 100 000 строк и 256 столбцов на лист.")
                rows = []
                for row in sheet.iter_rows():
                    rows.append(tuple(
                        UnsafeCell("Формула: замените её проверенным значением в копии исходного файла.")
                        if c.data_type == "f" else UnsafeCell("Ячейка содержит ошибку Excel.")
                        if c.data_type == "e" else c.value for c in row))
                    if len(rows) > MAX_ROWS or len(row) > MAX_COLUMNS:
                        raise ValueError("Лимит прототипа: 100 000 строк и 256 столбцов на лист.")
                sheets[sheet.title] = rows
        finally:
            workbook.close()
    if not sheets or not any(any(not blank(v) for v in row) for rows in sheets.values() for row in rows):
        raise ValueError("Файл не содержит данных.")
    return Source(path, hashlib.sha256(raw).hexdigest(), sheets)


def headers(rows: list[tuple], header_row: int) -> list[str]:
    if not 1 <= header_row <= len(rows):
        raise ValueError("Номер строки заголовков вне диапазона.")
    width = max((len(r) for r in rows), default=0)
    row = rows[header_row - 1]
    return [str(row[i]).strip() if i < len(row) and not blank(row[i]) else "" for i in range(width)]


def suggest_mapping(names: list[str], profile: dict) -> dict[str, int | None]:
    mapping = {}
    saved = profile.get("mapping", {})
    for f in profile["fields"]:
        if f["id"] in saved:
            matches = [i for i, name in enumerate(names) if name and name == saved[f["id"]]]
        else:
            aliases = {normalize(v) for v in [f["title"], *f.get("aliases", [])]}
            matches = [i for i, name in enumerate(names) if name and normalize(name) in aliases]
        mapping[f["id"]] = matches[0] if len(matches) == 1 else None
    return mapping


def suggest_header(rows: list[tuple], profile: dict) -> int | None:
    scores = [(sum(v is not None for v in suggest_mapping([str(x or "") for x in row], profile).values()), index)
              for index, row in enumerate(rows[:30], 1)]
    if not scores:
        return None
    best = max(score for score, _ in scores)
    winners = [index for score, index in scores if score == best]
    return winners[0] if best >= 2 and len(winners) == 1 else None


def parse_value(value: Any, spec: dict):
    if isinstance(value, UnsafeCell):
        raise ValueError(value.reason)
    if blank(value):
        if spec["required"]:
            raise ValueError("Обязательное значение отсутствует.")
        return None
    kind = spec["kind"]
    if isinstance(value, bool):
        raise ValueError("Логическое значение не поддерживается в этом поле.")
    if kind == "date":
        return flexible_date(value, spec)
    if kind == "identifier":
        if isinstance(value, (int, float)):
            if not math.isfinite(value) or value != int(value):
                raise ValueError("Идентификатор должен состоять из 12 цифр.")
            text = str(int(value))
        else:
            text = str(value).strip()
            if re.fullmatch(r'[0-9]{3}(?:[ -][0-9]{3}){3}', text):
                text = text.replace(' ', '').replace('-', '')
        if not re.fullmatch(r"[0-9]{12}", text):
            raise ValueError("Нужно ровно 12 цифр. Утраченные ведущие нули не восстанавливаются автоматически.")
        return text
    if kind == "amount":
        return flexible_amount(value, spec)
    if kind == "currency":
        text = str(value).strip().upper()
        aliases = {
            "KZT": ["ТЕНГЕ", "ТГ", "ТГ.", "₸", "ТЕҢГЕ", "КАЗАХСТАНСКИЙ ТЕНГЕ", "398"],
            "USD": ["ДОЛЛАР", "ДОЛЛАРЫ", "ДОЛЛАР США", "ДОЛЛАРЫ США", "US$", "840"],
            "EUR": ["ЕВРО", "€", "978"],
            "RUB": ["РУБ", "РУБ.", "РУБЛЬ", "РУБЛИ", "РОССИЙСКИЙ РУБЛЬ", "₽", "643"],
            "GBP": ["ФУНТ СТЕРЛИНГОВ", "ФУНТЫ СТЕРЛИНГОВ", "£", "826"],
            "CNY": ["ЮАНЬ", "ЮАНИ", "КИТАЙСКИЙ ЮАНЬ", "156"],
        }
        text = " ".join(text.split())
        for code, variants in aliases.items():
            if text in variants:
                return code
        if not re.fullmatch(r"[A-Z]{3}", text):
            raise ValueError("Валюта не распознана или неоднозначна (например $ или ¥). Укажите код: KZT, USD, EUR и т. д.")
        return text
    if not isinstance(value, str):
        if isinstance(value, (int, float, Decimal)) and math.isfinite(value):
            numeric = Decimal(str(value))
            if len(numeric.as_tuple().digits) > 15:
                raise ValueError('Числовой текст превышает точность Excel. Укажите исходный код вручную.')
            value = format(numeric, 'f')
            if '.' in value:
                value = value.rstrip('0').rstrip('.')
        else:
            raise ValueError('Значение нельзя преобразовать в текст. Укажите текст вручную.')
    text = value.strip()
    if len(text) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text):
        raise ValueError("Текст содержит недопустимые символы или превышает лимит Excel.")
    return text


def convert(source: Source, sheet: str, header_row: int, profile: dict, mapping: dict, corrections: dict | None = None) -> Result:
    validate_profile(profile)
    rows = source.sheets[sheet]
    names = headers(rows, header_row)
    result = Result(source, sheet, header_row, profile, dict(mapping))
    result.corrections = dict(corrections or {})
    used = []
    for f in profile["fields"]:
        index = mapping.get(f["id"])
        if index is None:
            if f["required"]:
                result.issues.append(Issue("error", header_row, f["title"], "Не выбран исходный столбец."))
        elif type(index) is not int or not 0 <= index < len(names):
            result.issues.append(Issue("error", header_row, f["title"], "Некорректный индекс столбца."))
        else:
            used.append(index)
    if len(used) != len(set(used)):
        result.issues.append(Issue("error", header_row, "Сопоставление", "Один исходный столбец выбран для нескольких полей."))
    if result.errors:
        return result
    ignored = [get_column_letter(i + 1) + ": " + (name or "без заголовка") for i, name in enumerate(names)
               if i not in used and any(i < len(row) and not blank(row[i]) for row in rows[header_row:])]
    if ignored:
        result.issues.append(Issue("warning", header_row, "Столбцы", "Не попадут в результат: " + "; ".join(ignored)))
    if header_row > 1:
        result.issues.append(Issue("warning", None, "Заголовок", f"Строки 1–{header_row - 1} находятся до выбранного заголовка и не обрабатываются."))
    if len(source.sheets) > 1:
        result.issues.append(Issue("warning", None, "Листы", "Обрабатывается только выбранный лист: " + sheet))
    seen = {}
    for number, row in enumerate(rows[header_row:], header_row + 1):
        if all(blank(value) for value in row):
            result.blank_rows.append(number)
            continue
        result.read_rows += 1
        record, bad = {}, False
        for f in profile["fields"]:
            index = mapping.get(f["id"])
            value = row[index] if index is not None and index < len(row) else None
            original = value
            key = (number, f["id"])
            if key in result.corrections:
                value = result.corrections[key]
                result.changes.append({"row": number, "field": f["id"], "kind": "manual", "before": str(original), "after": str(value)})
            try:
                record[f["id"]] = parse_value(value, f)
                if value is not None and str(value) != str(record[f["id"]]):
                    result.changes.append({"row": number, "field": f["id"], "kind": "normalization", "before": str(value), "after": str(record[f["id"]])})
            except ValueError as exc:
                result.issues.append(Issue("error", number, f["title"], str(exc), f["id"]))
                bad = True
        # A currency label inside an amount must agree with the dedicated currency field.
        currencies = [record.get(f['id']) for f in profile['fields'] if f['kind'] == 'currency' and record.get(f['id'])]
        for f in profile['fields']:
            if f['kind'] != 'amount' or f['id'] not in record:
                continue
            index = mapping.get(f['id'])
            raw_value = result.corrections.get((number, f['id']), row[index] if index is not None and index < len(row) else None)
            if raw_value is None:
                continue
            _, embedded_currency = amount_parts(raw_value)
            if embedded_currency and (not currencies or any(c != embedded_currency for c in currencies)):
                result.issues.append(Issue('error', number, f['title'], f'В сумме указана {embedded_currency}, но поле валюты отсутствует или отличается. Уточните валюту и сумму.', f['id']))
                bad = True
        if bad:
            result.invalid_rows.append(number)
        else:
            signature = tuple(record.values())
            if signature in seen:
                result.issues.append(Issue("warning", number, "Дубликат", f"Совпадает с преобразованной строкой {seen[signature]}. Обе строки сохранены."))
            else:
                seen[signature] = number
            result.records.append(record)
            result.source_rows.append(number)
    if not result.read_rows:
        result.issues.append(Issue("error", None, "Таблица", "После заголовка нет записей."))
    return result


def report_dict(result: Result) -> dict:
    return {
        "application_version": "0.1.0",
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "status": "blocked" if result.errors else "ready",
        "source_file": result.source.path.name,
        "source_sha256": result.source.sha256,
        "sheet": result.sheet,
        "header_row": result.header_row,
        "profile": result.profile,
        "mapping_indices_zero_based": result.mapping,
        "read_rows": result.read_rows,
        "valid_rows": len(result.records),
        "invalid_source_rows": result.invalid_rows,
        "blank_source_rows": result.blank_rows,
        "output_row_to_source_row": {str(i): row for i, row in enumerate(result.source_rows, 2)},
        "issues": [asdict(issue) for issue in result.issues],
        "changes": result.changes,
    }


def write_report(result: Result, path: Path):
    with Path(path).open("x", encoding="utf-8") as stream:
        json.dump(report_dict(result), stream, ensure_ascii=False, indent=2)


def export_result(result: Result, path: str | Path) -> tuple[Path, Path]:
    if result.errors or not result.records:
        raise ValueError("Выгрузка заблокирована: сначала исправьте ошибки и повторите проверку.")
    path = Path(path).resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Результат должен иметь расширение .xlsx.")
    report_path = path.with_suffix(".report.json")
    if path == result.source.path or path.exists() or report_path.exists():
        raise ValueError("Перезапись файлов запрещена. Выберите новое имя результата.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = result.profile["sheet_name"]
    fields = result.profile["fields"]
    for col, spec in enumerate(fields, 1):
        cell = sheet.cell(1, col, spec["title"])
        cell.data_type = "s"
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173E58")
        sheet.column_dimensions[get_column_letter(col)].width = min(40, max(18, len(spec["title"]) + 4))
    for row, record in enumerate(result.records, 2):
        for col, spec in enumerate(fields, 1):
            value = record[spec["id"]]
            cell = sheet.cell(row, col, value)
            if isinstance(value, str):
                cell.data_type = "s"  # Literal strings starting with '=' must never become formulas.
            if spec["kind"] == "identifier":
                cell.number_format = "@"
            elif spec["kind"] == "amount":
                cell.number_format = "#,##0.00"
            elif spec["kind"] == "date":
                cell.number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    # Stage serialization first; exclusive creation protects existing files even in a race.
    with tempfile.TemporaryDirectory(dir=path.parent) as temp:
        staged = Path(temp) / "result.xlsx"
        workbook.save(staged)
        workbook.close()
        made_output = made_report = False
        try:
            with path.open("xb") as stream:
                made_output = True
                stream.write(staged.read_bytes())
            with report_path.open("x", encoding="utf-8") as stream:
                made_report = True
                report = report_dict(result)
                report["status"] = "exported"
                report["output_file"] = path.name
                report["output_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
                json.dump(report, stream, ensure_ascii=False, indent=2)
        except Exception:
            if made_output:
                path.unlink(missing_ok=True)
            if made_report:
                report_path.unlink(missing_ok=True)
            raise
    return path, report_path
