"""Reusable synthetic workbooks for demonstrations and conversion regression checks."""
import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ['Дата', 'Документ', 'БИН', 'Наименование', 'Сумма', 'Валюта', 'Назначение']


def rows_for(values, column):
    rows = []
    for index, value in enumerate(values, 1):
        row = ['2026-10-06', f'{index:05d}', f'{index:012d}', f'Учебный контрагент {index}', '1250,50', 'KZT', 'Синтетическая операция']
        row[column] = value
        rows.append(row)
    return rows


def definitions():
    basic = rows_for(['1250,50', '2500,00', '-100,00'], 4)
    cases = []
    def add(name, title, rows, *, category='01_Автоматически', status='ready', header=1, valid=None, settings=None, note='', **extra):
        cases.append(dict(file=f'{category}/{name}', title=title, rows=rows, header_row=header,
                          sheet='Операции', status=status, valid_rows=len(rows) if valid is None else valid,
                          settings=settings or {}, note=note, **extra))
    add('01_обычный_реестр.xlsx', 'Обычный реестр', basic, note='Три записи; сумма в KZT: 3650,50.')
    add('02_разные_даты.xlsx', 'Даты в разных форматах', rows_for([
        '31/12/2026', '12/31/2026', '2026/12/31', '31-12-2026', '31 декабря 2026 г.',
        'December 31, 2026', '31 Dec 2026', datetime(2026, 12, 31), '31.12.2026 00:00:00'], 0),
        note='Все девять дат преобразуются в 2026-12-31.')
    add('03_разные_суммы.xlsx', 'Форматы сумм', rows_for([
        '1 234,56', '1,234.56', '1.234,56', '1’234.56', '1\u00a0234,56', '₸ 1 234,56',
        '1 234,56 KZT', '1234.5600', '1.23456e3'], 4), note='Все девять сумм преобразуются в 1234,56 KZT.')
    add('04_разные_валюты.xlsx', 'Названия и коды валют', rows_for(['тенге', '₸', '398', 'доллары США', '€', 'руб.', 'фунт стерлингов', 'юань'], 5),
        note='Ожидаются KZT, KZT, KZT, USD, EUR, RUB, GBP, CNY. Суммы по разным валютам не складываются.')
    add('05_возвраты_и_нулевые_суммы.xlsx', 'Возвраты и нули', rows_for(['(250,00)', '−250,00', '- 250.00', '0', '+250,00'], 4),
        note='Ожидаются -250, -250, -250, 0, 250; сумма в KZT: -500.')
    add('06_идентификаторы.xlsx', 'Ведущие нули и группировка', rows_for(['001234567890', '001 234 567 890', '001-234-567-890', 123456789012, '000000000001'], 2),
        note='Первые три идентификатора станут текстом 001234567890; нули сохраняются.')
    add('07_дубликаты_пустые_строки.xlsx', 'Дубликаты и пустые строки', [basic[0], [], basic[1], basic[0]], valid=3,
        note='Три записи сохранены, пустая строка пропущена. Дубликат отмечен предупреждением.')
    order = [3, 5, 4, 0, 2, 6, 1]
    add('08_переставленные_столбцы.xlsx', 'Другой порядок столбцов', [[r[i] for i in order] for r in basic],
        headers=['Получатель', 'Код валюты', 'Сумма операции', 'Дата платежа', 'БИН контрагента', 'Комментарий', '№ документа'],
        note='Поля определяются по синонимам, выходной порядок остаётся порядком профиля.')
    add('09_заголовок_после_шапки.xlsx', 'Заголовок после служебной шапки', basic, header=6,
        note='Строка заголовков — 6. Шапка объединена; строки до заголовка не выгружаются.')
    add('10_несколько_листов.xlsx', 'Несколько листов', basic, multiple=True,
        note='Выберите лист «Операции». Остальные листы не выгружаются; появится предупреждение.')
    add('11_необязательное_поле.xlsx', 'Нет назначения платежа', [r[:6] for r in basic], headers=HEADERS[:6],
        note='Назначение необязательно. Все три записи выгружаются с пустым назначением.')
    add('12_тысяча_операций.xlsx', 'Реестр на 1000 строк', rows_for(['100,00'] * 1000, 4),
        note='1000 записей, сумма в KZT: 100000,00. Предпросмотр покажет первые 200.')
    add('13_csv_точка_с_запятой.csv', 'CSV UTF-8 с разделителем ;', basic, delimiter=';', sheet_override='CSV',
        note='Три записи; дробные части записаны запятой.')
    add('14_csv_запятые_и_кавычки.csv', 'CSV с запятыми и кавычками', basic, delimiter=',', sheet_override='CSV',
        note='Значения с десятичной запятой заключены в кавычки; три записи читаются корректно.')
    category = '02_Нужны_настройки'
    add('15_порядок_даты.xlsx', 'Неоднозначный день и месяц', rows_for(['04/05/2026', '11/12/2026'], 0),
        category=category, status='settings', valid=0, settings={'date_order': 'dmy'},
        note='Авто: две ошибки. Выберите «День–месяц–год» → даты 2026-05-04 и 2026-12-11.')
    add('16_разделители_тысяч.xlsx', 'Неоднозначная запятая', rows_for(['1,234', '12,345', '1,234,567'], 4),
        category=category, status='settings', valid=1, settings={'decimal_separator': '.'},
        note='Авто: две ошибки. Выберите десятичный разделитель «Точка» → 1234, 12345, 1234567.')
    add('17_дата_со_временем.xlsx', 'Даты с ненулевым временем', rows_for([datetime(2026, 10, 6, 14, 30), '2026-10-07T09:15:00'], 0),
        category=category, status='settings', valid=0, settings={'drop_time': True},
        note='Авто: две ошибки. Включите «Отбрасывать время» → даты 2026-10-06 и 2026-10-07.')
    category = '03_Ошибки'
    add('18_несуществующие_даты.xlsx', 'Неверные календарные даты', rows_for(['31.02.2026', '2026-13-01', '31/12/26'], 0),
        category=category, status='blocked', valid=0, note='Три ошибки дат: несуществующий день, месяц и двузначный год. Исправляются в редакторе.')
    add('19_неверные_идентификаторы.xlsx', 'Неверные БИН/ИИН', rows_for(['12345', '', '00123456789X'], 2),
        category=category, status='blocked', valid=0, note='Три ошибки: короткий код, пустое обязательное поле и буква в коде.')
    add('20_неверные_суммы.xlsx', 'Ошибки сумм', rows_for(['1.005', '12 34,00', 'NaN'], 4),
        category=category, status='blocked', valid=0, note='Три ошибки: избыточная точность, неверная группировка тысяч, нечисловое значение.')
    add('21_неоднозначные_валюты.xlsx', 'Валюту нужно уточнить', rows_for(['$', '¥', 'неизвестно'], 5),
        category=category, status='blocked', valid=0, note='Три ошибки. В редакторе укажите выбранный трёхбуквенный код валюты.')
    add('22_конфликт_валют.xlsx', 'Валюта в сумме отличается', rows_for(['1250,50 EUR', 'USD 2500.00'], 4),
        category=category, status='blocked', valid=0, note='Две ошибки: в столбце валюты KZT, а в суммах EUR и USD. Нужно уточнить данные.')
    invalid = rows_for(['=100+200', '#DIV/0!'], 4)
    invalid.append(['ИТОГО', None, None, None, 300, None, None])
    add('23_формулы_ошибки_итоги.xlsx', 'Формулы, ошибка Excel и строка итога', invalid,
        category=category, status='blocked', valid=0,
        note='Формула и ошибка Excel блокируются. Строка итога не удаляется автоматически; отделите её от данных в рабочей копии.')
    add('24_неоднозначные_заголовки.xlsx', 'Дублирующиеся заголовки', [basic[0][:6]],
        headers=['Дата', 'Документ', 'БИН', 'ИИН', 'Сумма', 'Валюта'], category=category, status='blocked', valid=0,
        note='БИН и ИИН подходят одному полю; автоматический выбор запрещён. Также нет столбца контрагента: исправьте сопоставление и структуру.')
    return cases


def create_example_pack(directory: str | Path) -> list[Path]:
    directory = Path(directory)
    cases = definitions()
    paths = [directory / case['file'] for case in cases]
    metadata_paths = [directory / 'КАТАЛОГ.md', directory / 'examples.json']
    if any(path.exists() for path in [*paths, *metadata_paths]):
        raise ValueError('Файлы набора уже существуют. Выберите другую папку, чтобы сохранить исходные примеры.')
    metadata = []
    for case, path in zip(cases, paths):
        path.parent.mkdir(parents=True, exist_ok=True)
        labels = case.get('headers', HEADERS)
        if path.suffix == '.csv':
            with path.open('x', encoding='utf-8-sig', newline='') as stream:
                writer = csv.writer(stream, delimiter=case['delimiter'])
                writer.writerow(labels)
                writer.writerows(case['rows'])
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = case['sheet']
            if case.get('multiple'):
                wb.worksheets[0].title = 'Инструкция'
                ws = wb.create_sheet('Операции')
                wb.worksheets[0].append(['Синтетический пример. Выберите лист Операции.'])
                wb.create_sheet('Справка').append(['Этот лист не содержит операций.'])
            if case['header_row'] > 1:
                ws.append(['СИНТЕТИЧЕСКИЙ УЧЕБНЫЙ РЕЕСТР'])
                ws.merge_cells('A1:G1')
                ws.append(['Выгрузка учебной организации'])
                while ws.max_row < case['header_row'] - 1:
                    ws.append([None])
            ws.append(labels)
            for row in case['rows']:
                ws.append(row)
            for cell in ws[case['header_row']]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='173E58')
            for index, width in enumerate([28, 20, 24, 32, 25, 22, 32], 1):
                ws.column_dimensions[get_column_letter(index)].width = width
            ws.freeze_panes = f'A{case["header_row"] + 1}'
            ws.auto_filter.ref = f'A{case["header_row"]}:{get_column_letter(len(labels))}{ws.max_row}'
            with path.open('xb') as stream:
                wb.save(stream)
            wb.close()
        item = {key: case[key] for key in ('file', 'title', 'status', 'header_row', 'valid_rows', 'settings', 'note')}
        item['sheet'] = case.get('sheet_override', case['sheet'])
        item['records_after_settings'] = len(case['rows']) if case['status'] == 'settings' else case['valid_rows']
        metadata.append(item)
    with metadata_paths[1].open('x', encoding='utf-8') as stream:
        json.dump(metadata, stream, ensure_ascii=False, indent=2)
    guide = ['# Каталог учебных файлов', '', 'Все сведения синтетические. Номера БИН/ИИН не являются проверенными идентификаторами реальных организаций.', '',
             'Начните с папки **01_Автоматически**, затем попробуйте **02_Нужны_настройки** и **03_Ошибки**.', '',
             'Откройте файл → проверьте лист и заголовок → проверьте данные → посмотрите «Было → стало». Для каждого нового файла возвращайте настройки дат и разделителей в «Авто», а «Отбрасывать время» выключайте, если ниже не указано иное.', '',
             '| Файл | Что показать | Ожидаемый результат |', '| --- | --- | --- |']
    for item in metadata:
        guide.append(f'| [{Path(item["file"]).name}]({item["file"]}) | {item["title"]} | {item["note"]} |')
    guide.extend(['', 'Индексы и ожидаемые количества корректных записей доступны в `examples.json`. Выгрузка блокируется при любой ошибке; наличие корректных строк не разрешает частичную выгрузку.', ''])
    metadata_paths[0].write_text('\n'.join(guide), encoding='utf-8')
    return paths
