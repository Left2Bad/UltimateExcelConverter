"""Tk desktop UI. All file operations run locally."""

from __future__ import annotations

import copy
import json
import queue
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl.utils import get_column_letter

from .core import (convert, export_result, headers, load_profile, load_source,
                   suggest_header, suggest_mapping, write_report)
from .demo import create_demo


class Application(ttk.Frame):
    def __init__(self, root):
        super().__init__(root, padding=18)
        self.root = root
        self.pack(fill="both", expand=True)
        self.source = None
        self.result = None
        self.profile = load_profile()
        self.mapping_vars = {}
        self.mapping_options = []
        self.busy = False
        self.events = queue.Queue()
        self.sheet_var = tk.StringVar()
        self.header_var = tk.StringVar(value="1")
        self.status = tk.StringVar(value="Откройте файл или создайте учебные примеры.")
        self.filename = tk.StringVar(value="Файл не выбран")
        self.profile_label = tk.StringVar()
        self.build_ui()
        self.refresh_profile_label()
        self.header_var.trace_add("write", lambda *_: self.invalidate(clear_mapping=True))
        self.root.after(100, self.poll)

    def build_ui(self):
        self.root.title("Excel Converter · локальный прототип")
        self.root.geometry("1180x820")
        self.root.minsize(950, 660)
        style = ttk.Style()
        if "clam" in style.theme_names():
            style.theme_use("clam")
        style.configure("TFrame", background="#f4f7fa")
        style.configure("TLabel", background="#f4f7fa", foreground="#173e58", font=("Segoe UI", 10))
        style.configure("Title.TLabel", font=("Segoe UI", 23, "bold"))
        style.configure("TButton", font=("Segoe UI", 10), padding=(10, 6))
        style.configure("Treeview", rowheight=27, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        ttk.Label(self, text="Excel Converter", style="Title.TLabel").pack(anchor="w")
        ttk.Label(self, text="Подготовка таблиц • локальная обработка • проверка перед выгрузкой").pack(anchor="w", pady=(0, 12))
        ttk.Label(self, textvariable=self.profile_label, wraplength=1100).pack(anchor="w", pady=(0, 12))
        bar = ttk.Frame(self)
        bar.pack(fill="x")
        self.open_button = ttk.Button(bar, text="Открыть Excel / CSV", command=self.open_file)
        self.open_button.pack(side="left")
        self.demo_button = ttk.Button(bar, text="Создать демо", command=self.make_demo)
        self.demo_button.pack(side="left", padx=6)
        self.profile_button = ttk.Button(bar, text="Загрузить профиль", command=self.open_profile)
        self.profile_button.pack(side="left")
        self.save_profile_button = ttk.Button(bar, text="Сохранить профиль", command=self.save_profile)
        self.save_profile_button.pack(side="left", padx=6)
        ttk.Label(self, textvariable=self.filename, wraplength=1100).pack(anchor="w", pady=10)
        controls = ttk.Frame(self)
        controls.pack(fill="x", pady=(0, 12))
        ttk.Label(controls, text="Лист:").pack(side="left")
        self.sheet_box = ttk.Combobox(controls, textvariable=self.sheet_var, state="readonly", width=26)
        self.sheet_box.pack(side="left", padx=7)
        self.sheet_box.bind("<<ComboboxSelected>>", lambda _: self.select_sheet())
        ttk.Label(controls, text="Строка заголовков:").pack(side="left", padx=(12, 4))
        self.header_entry = ttk.Entry(controls, textvariable=self.header_var, width=7)
        self.header_entry.pack(side="left", padx=5)
        self.mapping_button = ttk.Button(controls, text="Применить заголовки", command=self.apply_headers)
        self.mapping_button.pack(side="left", padx=8)
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)
        self.source_tab = ttk.Frame(self.notebook, padding=10)
        self.mapping_tab = ttk.Frame(self.notebook, padding=10)
        self.result_tab = ttk.Frame(self.notebook, padding=10)
        self.issues_tab = ttk.Frame(self.notebook, padding=10)
        for frame, title in [(self.source_tab, "1. Исходные данные"), (self.mapping_tab, "2. Сопоставление"), (self.result_tab, "3. Предпросмотр"), (self.issues_tab, "4. Проверки")]:
            self.notebook.add(frame, text=title)
        ttk.Label(self.source_tab, text="Первые 200 строк выбранного листа. Номера соответствуют исходному файлу.").pack(anchor="w", pady=(0, 8))
        self.source_tree = self.make_tree(self.source_tab)
        ttk.Label(self.mapping_tab, text="Выберите источник каждого поля. * — обязательное поле. Сохранённый профиль запоминает названия столбцов.", wraplength=1050).pack(anchor="w", pady=(0, 10))
        # Scrollable mapping supports custom profiles with more fields than the demo.
        canvas_frame = ttk.Frame(self.mapping_tab)
        canvas_frame.pack(fill="both", expand=True)
        self.mapping_canvas = tk.Canvas(canvas_frame, highlightthickness=0, background="#f4f7fa")
        scroll = ttk.Scrollbar(canvas_frame, orient="vertical", command=self.mapping_canvas.yview)
        self.mapping_canvas.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        self.mapping_canvas.pack(side="left", fill="both", expand=True)
        self.mapping_frame = ttk.Frame(self.mapping_canvas)
        self.mapping_canvas.create_window((0, 0), window=self.mapping_frame, anchor="nw")
        self.mapping_frame.bind("<Configure>", lambda _: self.mapping_canvas.configure(scrollregion=self.mapping_canvas.bbox("all")))
        self.summary = tk.StringVar(value="Сначала выполните проверку.")
        ttk.Label(self.result_tab, textvariable=self.summary, wraplength=1050).pack(anchor="w", pady=(0, 8))
        self.result_tree = self.make_tree(self.result_tab)
        ttk.Label(self.issues_tab, text="Ошибки блокируют всю выгрузку. Предупреждения требуют просмотра. Двойной щелчок открывает сообщение целиком.", wraplength=1050).pack(anchor="w", pady=(0, 8))
        self.issues_tree = self.make_tree(self.issues_tab)
        self.issues_tree.bind("<Double-1>", self.show_issue)
        footer = ttk.Frame(self)
        footer.pack(fill="x", pady=(12, 0))
        self.check_button = ttk.Button(footer, text="Проверить и преобразовать", command=self.check)
        self.check_button.pack(side="left")
        self.export_button = ttk.Button(footer, text="Сохранить результат", command=self.export, state="disabled")
        self.export_button.pack(side="left", padx=8)
        self.report_button = ttk.Button(footer, text="Сохранить отчёт", command=self.save_report, state="disabled")
        self.report_button.pack(side="left")
        ttk.Label(self, textvariable=self.status, wraplength=1100).pack(anchor="w", pady=(10, 0))

    @staticmethod
    def make_tree(parent):
        frame = ttk.Frame(parent)
        frame.pack(fill="both", expand=True)
        tree = ttk.Treeview(frame, show="headings")
        vertical = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        horizontal = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        return tree

    @staticmethod
    def fill_tree(tree, columns, rows):
        tree.delete(*tree.get_children())
        ids = [str(i) for i in range(len(columns))]
        tree.configure(columns=ids)
        for key, title in zip(ids, columns):
            tree.heading(key, text=title)
            tree.column(key, width=175 if key != "0" else 95, minwidth=65, stretch=False)
        for row in rows:
            tree.insert("", "end", values=["" if v is None else str(v) for v in row])

    def refresh_profile_label(self):
        self.profile_label.set(self.profile["name"] + " · " + self.profile.get("description", "Пользовательский профиль"))

    def invalidate(self, clear_mapping=False):
        self.result = None
        self.export_button.configure(state="disabled")
        self.report_button.configure(state="disabled")
        self.result_tree.delete(*self.result_tree.get_children())
        self.issues_tree.delete(*self.issues_tree.get_children())
        self.summary.set("Настройки изменены. Выполните проверку заново.")
        if clear_mapping:
            self.mapping_vars.clear()
            for widget in self.mapping_frame.winfo_children():
                widget.destroy()

    def set_busy(self, busy):
        self.busy = busy
        for button in (self.open_button, self.demo_button, self.profile_button, self.save_profile_button, self.mapping_button, self.check_button):
            button.configure(state="disabled" if busy else "normal")
        self.sheet_box.configure(state="disabled" if busy else "readonly")
        self.header_entry.configure(state="disabled" if busy else "normal")
        for widget in self.mapping_frame.winfo_children():
            if isinstance(widget, ttk.Combobox):
                widget.configure(state="disabled" if busy else "readonly")
        self.export_button.configure(state="normal" if not busy and self.result and not self.result.errors else "disabled")
        self.report_button.configure(state="normal" if not busy and self.result else "disabled")

    def run_background(self, operation, on_success, text):
        if self.busy:
            return
        self.set_busy(True)
        self.status.set(text)
        def work():
            try:
                self.events.put((True, operation(), on_success))
            except Exception as exc:
                self.events.put((False, str(exc), on_success))
        threading.Thread(target=work, daemon=True).start()

    def poll(self):
        try:
            success, value, callback = self.events.get_nowait()
            self.set_busy(False)
            if success:
                callback(value)
            else:
                self.status.set("Операция не выполнена.")
                messagebox.showerror("Ошибка", value)
        except queue.Empty:
            pass
        self.root.after(100, self.poll)

    def open_file(self):
        path = filedialog.askopenfilename(filetypes=[("Excel / CSV", "*.xlsx *.csv"), ("Все файлы", "*.*")])
        if path:
            self.load_file(path)

    def load_file(self, path):
        self.run_background(lambda: load_source(path), self.loaded, "Чтение файла…")

    def loaded(self, source):
        self.source = source
        self.invalidate(clear_mapping=True)
        self.filename.set(str(source.path))
        self.sheet_box.configure(values=list(source.sheets))
        self.sheet_var.set(next(iter(source.sheets)))
        self.select_sheet()
        self.notebook.select(self.source_tab)

    def select_sheet(self):
        if not self.source:
            return
        rows = self.source.sheets[self.sheet_var.get()]
        suggestion = suggest_header(rows, self.profile)
        self.header_var.set(str(suggestion or 1))
        width = max((len(row) for row in rows[:200]), default=0)
        self.fill_tree(self.source_tree, ["Строка"] + [get_column_letter(i + 1) for i in range(width)],
                       [[i, *row] for i, row in enumerate(rows[:200], 1)])
        if rows:
            self.apply_headers()
        self.status.set(f"Лист «{self.sheet_var.get()}»: {len(rows)} строк. " +
                        (f"Предложена строка заголовков {suggestion}; проверьте её." if suggestion else "Заголовок не определён однозначно. Укажите строку вручную."))

    def apply_headers(self):
        if not self.source:
            return
        try:
            names = headers(self.source.sheets[self.sheet_var.get()], int(self.header_var.get()))
        except (ValueError, KeyError) as exc:
            messagebox.showerror("Заголовки", str(exc))
            return
        self.invalidate(clear_mapping=True)
        self.mapping_options = ["— не выбрано —"] + [f"{get_column_letter(i + 1)} · {name or '(без заголовка)'}" for i, name in enumerate(names)]
        suggestion = suggest_mapping(names, self.profile)
        for row, spec in enumerate(self.profile["fields"]):
            ttk.Label(self.mapping_frame, text=spec["title"] + (" *" if spec["required"] else ""), width=27).grid(row=row, column=0, sticky="w", pady=7)
            index = suggestion[spec["id"]]
            variable = tk.StringVar(value=self.mapping_options[index + 1 if index is not None else 0])
            self.mapping_vars[spec["id"]] = variable
            box = ttk.Combobox(self.mapping_frame, values=self.mapping_options, textvariable=variable, state="readonly", width=56)
            box.grid(row=row, column=1, sticky="w", padx=12)
            box.bind("<<ComboboxSelected>>", lambda _: self.invalidate())
        self.status.set("Сопоставление предложено. Проверьте его перед преобразованием.")

    def get_mapping(self):
        if not self.source or not self.mapping_vars:
            raise ValueError("Откройте файл и примените строку заголовков.")
        return {key: (self.mapping_options.index(var.get()) - 1 if self.mapping_options.index(var.get()) else None)
                for key, var in self.mapping_vars.items()}

    def check(self):
        try:
            mapping = self.get_mapping()
            header = int(self.header_var.get())
        except ValueError as exc:
            messagebox.showerror("Проверка", str(exc))
            return
        self.invalidate()
        source, sheet, profile = self.source, self.sheet_var.get(), copy.deepcopy(self.profile)
        self.run_background(lambda: convert(source, sheet, header, profile, mapping), self.checked, "Проверка и преобразование…")

    def checked(self, result):
        self.result = result
        warnings = len(result.issues) - len(result.errors)
        self.summary.set(f"Прочитано записей: {result.read_rows} • Корректных: {len(result.records)} • Ошибочных строк: {len(result.invalid_rows)} • Пустых: {len(result.blank_rows)}\n"
                         f"Ошибок: {len(result.errors)} • Предупреждений: {warnings}. Предпросмотр первых 200 корректных записей.")
        self.fill_tree(self.result_tree, ["Исх. строка"] + [f["title"] for f in result.profile["fields"]],
                       [[number, *record.values()] for number, record in list(zip(result.source_rows, result.records))[:200]])
        self.fill_tree(self.issues_tree, ["Уровень", "Строка", "Поле", "Описание"],
                       [["Ошибка" if issue.severity == "error" else "Внимание", issue.row, issue.field, issue.message] for issue in result.issues[:2000]])
        self.issues_tree.column("3", width=740)
        self.set_busy(False)
        self.status.set(("Выгрузка заблокирована. Исправьте данные в копии файла и откройте её заново." if result.errors else "Проверка завершена. Результат готов к сохранению.") +
                        (" Показаны первые 2000 замечаний; полный список — в отчёте." if len(result.issues) > 2000 else ""))
        self.notebook.select(self.issues_tab if result.errors else self.result_tab)

    def show_issue(self, _event):
        selection = self.issues_tree.selection()
        if selection:
            messagebox.showinfo("Замечание", "\n".join(str(v) for v in self.issues_tree.item(selection[0], "values")))

    def export(self):
        if not self.result or self.result.errors:
            return
        warnings = [i for i in self.result.issues if i.severity == "warning"]
        if warnings and not messagebox.askyesno("Проверьте предупреждения", f"Есть предупреждения: {len(warnings)}. Они доступны на вкладке «Проверки».\n\nВы просмотрели их и хотите сохранить результат?"):
            self.notebook.select(self.issues_tab)
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile=self.source.path.stem + "_converted.xlsx", filetypes=[("Excel", "*.xlsx")])
        if path:
            result = self.result
            self.run_background(lambda: export_result(result, path), self.exported, "Сохранение результата и отчёта…")

    def exported(self, paths):
        self.status.set(f"Сохранено: {paths[0]}")
        messagebox.showinfo("Готово", f"Результат:\n{paths[0]}\n\nОтчёт:\n{paths[1]}")

    def save_report(self):
        if not self.result:
            return
        path = filedialog.asksaveasfilename(defaultextension=".json", initialfile="validation.report.json", filetypes=[("JSON", "*.json")])
        if path:
            try:
                write_report(self.result, Path(path))
                self.status.set(f"Отчёт сохранён: {path}")
            except Exception as exc:
                messagebox.showerror("Отчёт", str(exc))

    def open_profile(self):
        path = filedialog.askopenfilename(filetypes=[("JSON-профиль", "*.json")])
        if not path:
            return
        try:
            profile = load_profile(path)
        except Exception as exc:
            messagebox.showerror("Профиль", str(exc))
            return
        self.profile = profile
        self.refresh_profile_label()
        self.invalidate(clear_mapping=True)
        if self.source:
            self.select_sheet()

    def save_profile(self):
        try:
            mapping = self.get_mapping()
            names = headers(self.source.sheets[self.sheet_var.get()], int(self.header_var.get()))
            if any(index is not None and (not names[index] or names.count(names[index]) != 1) for index in mapping.values()):
                raise ValueError("Для сохранения профиля выбранные заголовки должны быть непустыми и уникальными.")
            profile = copy.deepcopy(self.profile)
            profile["mapping"] = {key: names[index] if index is not None else None for key, index in mapping.items()}
            path = filedialog.asksaveasfilename(defaultextension=".json", initialfile="custom_profile.json", filetypes=[("JSON-профиль", "*.json")])
            if path:
                with Path(path).open("x", encoding="utf-8") as stream:
                    json.dump(profile, stream, ensure_ascii=False, indent=2)
                self.status.set(f"Профиль сохранён: {path}")
        except Exception as exc:
            messagebox.showerror("Профиль", str(exc))

    def make_demo(self):
        directory = filedialog.askdirectory(title="Папка для синтетических примеров")
        if directory:
            self.run_background(lambda: create_demo(directory), lambda paths: self.load_file(paths[0]), "Создание учебных примеров…")


def launch():
    root = tk.Tk()
    Application(root)
    root.mainloop()
