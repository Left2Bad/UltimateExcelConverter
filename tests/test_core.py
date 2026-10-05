import copy
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_converter.core import (Source, UnsafeCell, convert, export_result, headers,
                                  load_profile, load_source, parse_value, suggest_header,
                                  suggest_mapping, validate_profile)
from excel_converter.demo import create_demo


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.profile = load_profile()

    def spec(self, kind):
        return next(f for f in self.profile["fields"] if f["kind"] == kind)

    def converted(self, path):
        source = load_source(path)
        sheet = next(iter(source.sheets))
        row = suggest_header(source.sheets[sheet], self.profile)
        mapping = suggest_mapping(headers(source.sheets[sheet], row), self.profile)
        return convert(source, sheet, row, self.profile, mapping)

    def test_roundtrip_and_report(self):
        good, _, _ = create_demo(self.directory)
        before = good.read_bytes()
        result = self.converted(good)
        self.assertFalse(result.errors)
        self.assertEqual(result.header_row, 3)
        self.assertEqual(result.source_rows, [4, 5, 6])
        self.assertEqual(sum(r['amount'] for r in result.records), Decimal('129250.50'))
        output, report = export_result(result, self.directory / 'output.xlsx')
        wb = load_workbook(output)
        self.addCleanup(wb.close)
        ws = wb.active
        self.assertEqual(ws.title, 'Операции')
        self.assertEqual(ws.max_row, 4)
        self.assertEqual([c.value for c in ws[1]], [f['title'] for f in self.profile['fields']])
        self.assertEqual(ws['C2'].value, '001234567890')
        self.assertEqual(ws['C2'].data_type, 's')
        self.assertEqual(ws['E2'].value, 125000.5)
        self.assertEqual(ws['A2'].value, datetime(2026, 10, 1))
        audit = json.loads(report.read_text(encoding='utf-8'))
        self.assertEqual(audit['output_row_to_source_row'], {'2': 4, '3': 5, '4': 6})
        self.assertEqual(audit['status'], 'exported')
        self.assertEqual(good.read_bytes(), before)

    def test_invalid_rows_block_whole_export(self):
        _, bad, _ = create_demo(self.directory)
        result = self.converted(bad)
        self.assertEqual(result.invalid_rows, [7, 8, 9])
        self.assertEqual(result.read_rows, 6)
        self.assertTrue(any('Формула' in i.message for i in result.errors))
        target = self.directory / 'blocked.xlsx'
        with self.assertRaises(ValueError):
            export_result(result, target)
        self.assertFalse(target.exists())

    def test_csv_reordered_columns(self):
        _, _, csv = create_demo(self.directory)
        result = self.converted(csv)
        self.assertFalse(result.errors)
        self.assertEqual(result.records[0]['identifier'], '001234567890')
        self.assertEqual(result.records[0]['document'], '0001')

    def test_ambiguous_headers_never_guessed(self):
        names = ['БИН', 'ИИН', 'Дата', 'Сумма']
        mapping = suggest_mapping(names, self.profile)
        self.assertIsNone(mapping['identifier'])
        self.assertIsNone(suggest_header([tuple(names), tuple(names)], self.profile))

    def test_saved_mapping_requires_exact_unique_name(self):
        self.profile['mapping'] = {'identifier': 'Регистрационный код', 'purpose': None}
        mapping = suggest_mapping(['Регистрационный код', 'Назначение'], self.profile)
        self.assertEqual(mapping['identifier'], 0)
        self.assertIsNone(mapping['purpose'])
        self.assertIsNone(suggest_mapping(['Регистрационный код'] * 2, self.profile)['identifier'])

    def test_identifier_never_padded(self):
        self.assertEqual(parse_value('001234567890', self.spec('identifier')), '001234567890')
        self.assertEqual(parse_value(123456789012, self.spec('identifier')), '123456789012')
        for value in [1234567890, 1.25, '12345678901', '１２３４５６７８９０１２', float('inf')]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_value(value, self.spec('identifier'))

    def test_strict_money(self):
        for value, expected in [('125 000,50', '125000.50'), ('1\u00a0234.20', '1234.20'), (0, '0'), ('-250,00', '-250.00')]:
            self.assertEqual(parse_value(value, self.spec('amount')), Decimal(expected))
        for value in ['1,234', '12 34', 'NaN', 'Infinity', '1000000000000000', '1.005']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_value(value, self.spec('amount'))

    def test_strict_dates(self):
        self.assertEqual(parse_value('01.10.2026', self.spec('date')), date(2026, 10, 1))
        self.assertEqual(parse_value('2026-10-01', self.spec('date')), date(2026, 10, 1))
        for value in ['01/02/2026', '31.02.2026', 45000, datetime(2026, 1, 1, 10, 30)]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_value(value, self.spec('date'))

    def test_flexible_amount_formats(self):
        for value, expected in [('1,234.56', '1234.56'), ('1.234,56', '1234.56'), ("1’234.56", '1234.56'), ('(250,00)', '-250'), ('− 250,00', '-250'), ('1,234,567', '1234567'), ('1e5', '100000'), ('1250.5000', '1250.5'), ('₸ 1 250,50', '1250.50'), ('1 250,50 тг.', '1250.50')]:
            with self.subTest(value=value):
                self.assertEqual(parse_value(value, self.spec('amount')), Decimal(expected))
        for value in ['1,23,456', '1.234.56', '(+250)', '123456789012345.000000000000000001', '12,3456', '12 345,67.8']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_value(value, self.spec('amount'))
        self.assertEqual(parse_value('1,234', dict(self.spec('amount'), decimal_separator='.')), Decimal(1234))
        with self.assertRaises(ValueError):
            parse_value('1,234', dict(self.spec('amount'), decimal_separator=','))

    def test_flexible_dates_and_explicit_policy(self):
        for value in ['31/12/2026', '12/31/2026', '31-12-2026', '2026/12/31', '31 декабря 2026 г.', '31 Dec 2026', 'December 31, 2026', '31.12.2026 00:00:00']:
            with self.subTest(value=value):
                self.assertEqual(parse_value(value, self.spec('date')), date(2026, 12, 31))
        self.assertEqual(parse_value('01/02/2026', dict(self.spec('date'), date_order='dmy')), date(2026, 2, 1))
        self.assertEqual(parse_value('01/02/2026', dict(self.spec('date'), date_order='mdy')), date(2026, 1, 2))
        self.assertEqual(parse_value('2026-12-31T12:30:00', dict(self.spec('date'), drop_time=True)), date(2026, 12, 31))
        for value in ['2026-12-31T25:00', '31/12/26', '2026-12-31T12:00:00Z']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_value(value, dict(self.spec('date'), drop_time=True))

    def test_embedded_currency_mismatch(self):
        good, _, _ = create_demo(self.directory)
        result = self.converted(good)
        fixed = convert(result.source, result.sheet, result.header_row, result.profile, result.mapping, {(4, 'amount'): '1 250,50 EUR'})
        self.assertIn(4, fixed.invalid_rows)
        fixed = convert(result.source, result.sheet, result.header_row, result.profile, result.mapping, {(4, 'amount'): '1 250,50 KZT'})
        self.assertFalse(fixed.errors)
        self.assertEqual(fixed.records[0]['amount'], Decimal('1250.50'))

    def test_currency_variants(self):
        for value, code in [('тенге', 'KZT'), (' тг. ', 'KZT'), ('₸', 'KZT'), ('398', 'KZT'), ('доллары США', 'USD'), ('€', 'EUR'), ('руб.', 'RUB')]:
            self.assertEqual(parse_value(value, self.spec('currency')), code)
        for value in ['$', '¥', 'неизвестно']:
            with self.assertRaises(ValueError):
                parse_value(value, self.spec('currency'))

    def test_manual_corrections_revalidate_and_preserve_source(self):
        _, bad, _ = create_demo(self.directory)
        before = bad.read_bytes()
        result = self.converted(bad)
        fixes = {(7, 'date'): '2026-05-04', (7, 'identifier'): '000000012345', (7, 'amount'): '1234', (8, 'amount'): '300'}
        corrected = convert(result.source, result.sheet, result.header_row, result.profile, result.mapping, fixes)
        self.assertEqual(corrected.invalid_rows, [9])
        self.assertEqual(corrected.records[3]['currency'], 'KZT')
        self.assertTrue(any(c['kind'] == 'manual' for c in corrected.changes))
        self.assertTrue(any(c['kind'] == 'normalization' for c in corrected.changes))
        self.assertEqual(bad.read_bytes(), before)
        fixes[(7, 'date')] = 'bad date'
        rejected = convert(result.source, result.sheet, result.header_row, result.profile, result.mapping, fixes)
        self.assertIn(7, rejected.invalid_rows)

    def test_numeric_document_to_text(self):
        self.assertEqual(parse_value(123, self.spec('text')), '123')
        self.assertEqual(parse_value(123.0, self.spec('text')), '123')
        self.assertEqual(parse_value('000123', self.spec('text')), '000123')
        self.assertEqual(parse_value(1234.56, dict(self.spec('amount'), decimal_separator=',')), Decimal('1234.56'))

    def test_optional_blank_and_boolean(self):
        optional = self.profile['fields'][-1]
        self.assertIsNone(parse_value(' ', optional))
        with self.assertRaises(ValueError):
            parse_value(True, self.spec('amount'))

    def test_duplicate_rows_preserved_and_blank_counted(self):
        good, _, _ = create_demo(self.directory)
        source = load_source(good)
        source.sheets['Реестр'].extend([tuple(), source.sheets['Реестр'][3]])
        mapping = suggest_mapping(headers(source.sheets['Реестр'], 3), self.profile)
        result = convert(source, 'Реестр', 3, self.profile, mapping)
        self.assertEqual(len(result.records), 4)
        self.assertEqual(result.blank_rows, [7])
        self.assertEqual(result.source_rows[-1], 8)
        self.assertTrue(any(i.field == 'Дубликат' for i in result.issues))

    def test_invalid_mapping_blocks(self):
        good, _, _ = create_demo(self.directory)
        source = load_source(good)
        mapping = suggest_mapping(headers(source.sheets['Реестр'], 3), self.profile)
        for bad_index in [None, 999, True, -1]:
            changed = dict(mapping, date=bad_index)
            self.assertTrue(convert(source, 'Реестр', 3, self.profile, changed).errors)
        mapping['party'] = mapping['document']
        self.assertTrue(convert(source, 'Реестр', 3, self.profile, mapping).errors)

    def test_no_overwrite_source_output_or_report(self):
        good, _, _ = create_demo(self.directory)
        result = self.converted(good)
        with self.assertRaises(ValueError):
            export_result(result, good)
        out = self.directory / 'converted.xlsx'
        export_result(result, out)
        before = out.read_bytes()
        with self.assertRaises(ValueError):
            export_result(result, out)
        self.assertEqual(out.read_bytes(), before)
        other = self.directory / 'other.report.json'
        other.write_text('keep', encoding='utf-8')
        with self.assertRaises(ValueError):
            export_result(result, self.directory / 'other.xlsx')
        self.assertEqual(other.read_text(), 'keep')

    def test_literal_formula_text_export(self):
        good, _, _ = create_demo(self.directory)
        result = self.converted(good)
        result.records[0]['purpose'] = '=HYPERLINK("https://invalid.example","demo")'
        out, _ = export_result(result, self.directory / 'text.xlsx')
        wb = load_workbook(out, data_only=False)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['G2'].data_type, 's')
        self.assertTrue(wb.active['G2'].value.startswith('='))

    def test_excel_error_and_formula_cells_rejected(self):
        for value in [UnsafeCell('Формула'), UnsafeCell('Ошибка Excel')]:
            with self.assertRaises(ValueError):
                parse_value(value, self.spec('amount'))

    def test_empty_file_and_old_format(self):
        empty = self.directory / 'empty.csv'
        empty.write_text('', encoding='utf-8')
        with self.assertRaises(ValueError):
            load_source(empty)
        with self.assertRaises(ValueError):
            load_source(self.directory / 'legacy.xls')

    def test_profile_validation(self):
        for mutate in [lambda p: p.update(version=2), lambda p: p.update(sheet_name='Bad/Name'),
                       lambda p: p['fields'].append(copy.deepcopy(p['fields'][0])),
                       lambda p: p['fields'][0].update(kind='guess'),
                       lambda p: p.update(mapping={'missing': 'Дата'})]:
            profile = copy.deepcopy(self.profile)
            mutate(profile)
            with self.assertRaises(ValueError):
                validate_profile(profile)

    def test_cli_requires_warning_acknowledgment(self):
        good, _, _ = create_demo(self.directory)
        output = self.directory / 'cli.xlsx'
        command = [sys.executable, '-m', 'excel_converter', 'convert', str(good), str(output)]
        run = subprocess.run(command, capture_output=True)
        self.assertEqual(run.returncode, 3, run.stderr)
        self.assertFalse(output.exists())
        run = subprocess.run(command + ['--accept-warnings'], capture_output=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertTrue(output.exists())


if __name__ == '__main__':
    unittest.main()
